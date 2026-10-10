"""Configuration, text extraction, upload validation and chunking (no models needed)."""

import io

import docx
import pytest
from pydantic import ValidationError

from adaptive_rag.config import Settings
from adaptive_rag.ingestion.chunking import chunk_pages, chunk_text
from adaptive_rag.ingestion.extract import (
    ExtractionError,
    PageText,
    extract_pages,
    file_type_for,
    safe_filename,
)
from tests.conftest import make_pdf


# ---------------------------------------------------------------- configuration
def test_dev_mode_resolves_lightweight_backends():
    s = Settings(mode="dev", _env_file=None)
    assert (s.resolved_embedding_backend, s.resolved_llm_backend, s.resolved_nli_backend) == ("hash", "extractive", "none")


def test_real_mode_resolves_local_models():
    s = Settings(mode="real", _env_file=None)
    assert (s.resolved_embedding_backend, s.resolved_llm_backend, s.resolved_nli_backend) == (
        "fastembed", "ollama", "transformers")


def test_explicit_backend_overrides_mode():
    assert Settings(mode="real", llm_backend="extractive", _env_file=None).resolved_llm_backend == "extractive"


@pytest.mark.parametrize("kwargs", [
    {"chunk_size": 300, "chunk_overlap": 300},
    {"accept_threshold": 0.3, "abstain_threshold": 0.5},
    {"max_retries": 3},
    {"w_relevance": 0, "w_coverage": 0, "w_verification": 0, "w_consistency": 0},
])
def test_invalid_settings_rejected(kwargs):
    with pytest.raises(ValidationError):
        Settings(_env_file=None, **kwargs)


def test_env_variables_are_read(monkeypatch):
    monkeypatch.setenv("ARAG_TOP_K", "7")
    monkeypatch.setenv("ARAG_ACCEPT_THRESHOLD", "0.8")
    s = Settings(_env_file=None)
    assert s.top_k == 7 and s.accept_threshold == 0.8


def test_public_dict_contains_resolved_backends_and_no_paths():
    d = Settings(mode="dev", _env_file=None).public_dict()
    assert d["resolved"]["llm_backend"] == "extractive"
    assert "runtime_dir" not in d and "data_dir" not in d


# ---------------------------------------------------------------- filenames / types
@pytest.mark.parametrize("raw,expected", [
    ("../../etc/passwd.txt", "passwd.txt"),
    ("C:\\Users\\x\\report.pdf", "report.pdf"),
    ("my notes (v2).md", "my notes (v2).md"),
    ("a<b>c|d?.txt", "a_b_c_d_.txt"),
    ("", "document"),
])
def test_safe_filename(raw, expected):
    assert safe_filename(raw) == expected


def test_file_type_validation():
    assert file_type_for("A.PDF") == "pdf"
    assert file_type_for("x.markdown") == "markdown"
    with pytest.raises(ExtractionError, match="Unsupported"):
        file_type_for("virus.exe")
    with pytest.raises(ExtractionError):
        file_type_for("noextension")


# ---------------------------------------------------------------- extraction
def test_extract_txt_and_markdown():
    assert extract_pages(b"Hello   world\n\n\n\nBye", "txt")[0].text == "Hello world\n\nBye"
    assert "# Title" in extract_pages(b"# Title\nBody", "markdown")[0].text


def test_extract_utf16_and_latin1_text():
    assert extract_pages("Caf\u00e9 notes".encode("utf-16"), "txt")[0].text == "Caf\u00e9 notes"
    assert extract_pages("Caf\u00e9".encode("latin-1"), "txt")[0].text == "Caf\u00e9"


def test_extract_docx_paragraphs_and_tables():
    d = docx.Document()
    d.add_paragraph("The reactor uses molten salt.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Coolant", "Salt"
    buf = io.BytesIO()
    d.save(buf)
    text = extract_pages(buf.getvalue(), "docx")[0].text
    assert "molten salt" in text and "Coolant | Salt" in text


def test_extract_pdf_keeps_page_numbers():
    pages = extract_pages(make_pdf(["First page about otters.", "Second page about beavers."]), "pdf")
    assert [p.page for p in pages] == [1, 2]
    assert "beavers" in pages[1].text


def test_pdf_without_text_layer_is_rejected():
    with pytest.raises(ExtractionError, match="OCR"):
        extract_pages(make_pdf([""]), "pdf")


@pytest.mark.parametrize("data,ftype", [(b"", "txt"), (b"   \n ", "txt"), (b"not a docx", "docx"), (b"%PDF-broken", "pdf")])
def test_bad_or_empty_files_raise(data, ftype):
    with pytest.raises(ExtractionError):
        extract_pages(data, ftype)


# ---------------------------------------------------------------- chunking
def test_short_text_is_single_chunk():
    assert chunk_text("Short text.", 200, 50) == [(0, 11)]


def test_chunks_cover_text_with_overlap_and_size_limit():
    text = " ".join(f"Sentence number {i} talks about topic {i % 7}." for i in range(200))
    spans = chunk_text(text, 400, 80)
    assert spans[0][0] == 0 and spans[-1][1] == len(text)
    for (s1, e1), (s2, _) in zip(spans, spans[1:]):
        assert e1 - s1 <= 400
        assert s2 < e1, "consecutive chunks must overlap"
        assert s2 > s1, "chunking must make progress"


def test_chunks_prefer_sentence_boundaries():
    text = " ".join(f"This is sentence {i} of the test." for i in range(60))
    for s, e in chunk_text(text, 300, 50)[:-1]:
        assert text[s:e].rstrip().endswith(".")


def test_chunk_ids_and_pages():
    chunks = chunk_pages([PageText(1, "Alpha. " * 100), PageText(2, "Beta. " * 100)], "doc1", "f.pdf", 250, 40)
    assert [c.chunk_index for c in chunks] == list(range(len(chunks)))
    assert all(c.chunk_id == f"doc1:{c.chunk_index}" for c in chunks)
    assert {c.page for c in chunks} == {1, 2}
    assert all(("Alpha" in c.text) != ("Beta" in c.text) for c in chunks), "chunks must not cross pages"
