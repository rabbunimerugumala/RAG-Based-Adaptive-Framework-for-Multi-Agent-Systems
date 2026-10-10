"""Text extraction for PDF (text layer only), DOCX, TXT and Markdown.

Scanned PDFs without a text layer are NOT supported (no OCR is implemented); they are rejected with
a clear error instead of being indexed as empty documents.
"""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass
from pathlib import PurePath

SUPPORTED_EXTENSIONS = {".pdf": "pdf", ".docx": "docx", ".txt": "txt", ".md": "markdown", ".markdown": "markdown"}


class ExtractionError(ValueError):
    """Raised when a file cannot be turned into indexable text."""


@dataclass
class PageText:
    page: int | None  # 1-based page number for PDFs, None otherwise
    text: str


def safe_filename(name: str) -> str:
    """Strip directories and unsafe characters; keep the extension. Never returns an empty name."""
    name = PurePath(name.replace("\\", "/")).name
    name = unicodedata.normalize("NFKC", name)
    name = re.sub(r"[^\w.\- ()]", "_", name).strip(" .")
    name = re.sub(r"_+", "_", name)
    return (name or "document")[:150]


def file_type_for(filename: str) -> str:
    ext = PurePath(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(
            f"Unsupported file type '{ext or '(none)'}'. Supported: PDF, DOCX, TXT, Markdown."
        )
    return SUPPORTED_EXTENSIONS[ext]


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # de-hyphenate line breaks
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _decode(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def extract_pages(data: bytes, file_type: str) -> list[PageText]:
    if not data:
        raise ExtractionError("The file is empty.")
    if file_type == "pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ExtractionError("Encrypted PDFs are not supported.")
            pages = [PageText(i + 1, _clean(p.extract_text() or "")) for i, p in enumerate(reader.pages)]
        except ExtractionError:
            raise
        except Exception as exc:  # pypdf raises many error types for corrupt files
            raise ExtractionError(f"Could not read PDF (corrupt or not a PDF): {exc}") from exc
        pages = [p for p in pages if p.text]
        if not pages:
            raise ExtractionError(
                "No extractable text found. Scanned/image-only PDFs need OCR, which is not implemented."
            )
        return pages
    if file_type == "docx":
        import docx

        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:  # corrupt/renamed files raise many different parser errors
            raise ExtractionError(f"Could not read DOCX (corrupt or not a Word document): {exc}") from exc
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = _clean("\n\n".join(parts))
    else:  # txt / markdown
        text = _clean(_decode(data))
    if not text:
        raise ExtractionError("The document contains no text.")
    return [PageText(None, text)]
